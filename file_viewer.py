import customtkinter as ctk
from tkinter import scrolledtext, messagebox
from PIL import Image, ImageTk
import io
import tempfile
import os
import threading
import time
import subprocess

try:
    import fitz  # PyMuPDF for PDF viewing
except ImportError:
    fitz = None

try:
    import cv2
except ImportError:
    cv2 = None

try:
    import pygame
except ImportError:
    pygame = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

class FileViewer(ctk.CTkToplevel):
    def __init__(self, master, file_bytes, filename, *args, **kwargs):
        super().__init__(master, *args, **kwargs)
        self.title(f"Enhanced File Viewer - {filename}")
        self.geometry("1000x700")
        self.minsize(800, 600)
        
        self.file_bytes = file_bytes
        self.filename = filename
        self.file_ext = os.path.splitext(filename)[1].lower()
        
        # Initialize variables for media control
        self.current_pdf_page = 0
        self.total_pdf_pages = 0
        self.pdf_doc = None
        self.video_cap = None
        self.video_playing = False
        self.video_paused = False
        self.video_thread = None
        self.audio_position = 0
        self.audio_playing = False
        self.audio_paused = False
        self.video_has_audio = False
        self.video_audio_path = None
        
        self.protocol("WM_DELETE_WINDOW", self.on_closing)
        
        # Create main UI
        self._create_ui()
        self._show_file()

    def _initialize_audio(self):
        """Initialize pygame mixer with proper error handling"""
        try:
            # Quit any existing mixer
            if pygame.mixer.get_init():
                pygame.mixer.quit()
            
            # Initialize with compatible settings
            pygame.mixer.pre_init(frequency=22050, size=-16, channels=2, buffer=512)
            pygame.mixer.init()
            return True
        except Exception as e:
            print(f"Audio initialization failed: {e}")
            return False

    def _show_audio_status(self, message, is_error=False):
        """Display audio status messages to user"""
        color = "red" if is_error else "green"
        if hasattr(self, 'audio_status_label'):
            self.audio_status_label.configure(text=message, text_color=color)
        print(f"Audio Status: {message}")

    def _create_ui(self):
        """Create the main UI layout"""
        # Main container
        self.main_frame = ctk.CTkFrame(self)
        self.main_frame.pack(fill="both", expand=True, padx=10, pady=10)
        
        # Header with file info
        self.header_frame = ctk.CTkFrame(self.main_frame)
        self.header_frame.pack(fill="x", padx=5, pady=(5, 10))
        
        self.file_info_label = ctk.CTkLabel(
            self.header_frame,
            text=f"File: {self.filename} | Size: {len(self.file_bytes)} bytes",
            font=ctk.CTkFont(size=12, weight="bold")
        )
        self.file_info_label.pack(pady=10)
        
        # Content area
        self.content_frame = ctk.CTkFrame(self.main_frame)
        self.content_frame.pack(fill="both", expand=True, padx=5, pady=5)
        
        # Control panel (will be populated based on file type)
        self.control_frame = ctk.CTkFrame(self.main_frame)
        self.control_frame.pack(fill="x", padx=5, pady=(5, 5))

    def _show_file(self):
        """Display file based on its type"""
        if self.file_ext in [".txt", ".py", ".md", ".log", ".csv", ".json", ".xml"]:
            self._show_text()
        elif self.file_ext in [".jpg", ".jpeg", ".png", ".bmp", ".gif", ".webp"]:
            self._show_image()
        elif self.file_ext in [".pdf"] and fitz:
            self._show_pdf()
        elif self.file_ext in [".mp4", ".avi", ".mov", ".mkv", ".wmv", ".flv"] and cv2:
            self._show_video()
        elif self.file_ext in [".mp3", ".wav", ".ogg", ".m4a", ".flac"] and pygame:
            self._show_audio()
        elif self.file_ext in [".pptx"] and Presentation:
            self._show_pptx()
        else:
            self._show_unsupported()

    def _show_text(self):
        """Enhanced text viewer with search and line numbers"""
        try:
            text = self.file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            try:
                text = self.file_bytes.decode('latin-1')
            except UnicodeDecodeError:
                text = self.file_bytes.decode('utf-8', errors='replace')
        
        # Search controls
        search_frame = ctk.CTkFrame(self.control_frame)
        search_frame.pack(fill="x", padx=5, pady=5)
        
        self.search_entry = ctk.CTkEntry(search_frame, placeholder_text="Search text...")
        self.search_entry.pack(side="left", padx=5, pady=5, fill="x", expand=True)
        
        search_btn = ctk.CTkButton(search_frame, text="Find", command=self._search_text, width=60)
        search_btn.pack(side="left", padx=5, pady=5)
        
        # Text display options
        options_frame = ctk.CTkFrame(self.control_frame)
        options_frame.pack(fill="x", padx=5, pady=5)
        
        self.line_numbers_var = ctk.BooleanVar(value=True)
        line_numbers_cb = ctk.CTkCheckBox(options_frame, text="Line Numbers",
                                        variable=self.line_numbers_var,
                                        command=self._toggle_line_numbers)
        line_numbers_cb.pack(side="left", padx=5, pady=5)
        
        self.word_wrap_var = ctk.BooleanVar(value=True)
        word_wrap_cb = ctk.CTkCheckBox(options_frame, text="Word Wrap",
                                     variable=self.word_wrap_var,
                                     command=self._toggle_word_wrap)
        word_wrap_cb.pack(side="left", padx=5, pady=5)
        
        # Font size controls
        ctk.CTkLabel(options_frame, text="Font Size:").pack(side="left", padx=(20, 5), pady=5)
        self.font_size_var = ctk.IntVar(value=12)
        font_slider = ctk.CTkSlider(options_frame, from_=8, to=24, variable=self.font_size_var,
                                  command=self._change_font_size, width=100)
        font_slider.pack(side="left", padx=5, pady=5)
        
        # Text widget
        self.text_widget = scrolledtext.ScrolledText(
            self.content_frame,
            wrap="word" if self.word_wrap_var.get() else "none",
            font=("Consolas", 12),
            state="normal"
        )
        self.text_widget.pack(expand=True, fill="both", padx=10, pady=10)
        
        # Add line numbers and text
        self._update_text_display(text)

    def _search_text(self):
        """Search functionality for text viewer"""
        search_term = self.search_entry.get()
        if search_term:
            self.text_widget.tag_remove("highlight", "1.0", "end")
            start_pos = "1.0"
            while True:
                pos = self.text_widget.search(search_term, start_pos, "end")
                if not pos:
                    break
                end_pos = f"{pos}+{len(search_term)}c"
                self.text_widget.tag_add("highlight", pos, end_pos)
                start_pos = end_pos
            self.text_widget.tag_config("highlight", background="yellow", foreground="black")

    def _toggle_line_numbers(self):
        """Toggle line numbers display"""
        self._update_text_display(self.text_widget.get("1.0", "end-1c"))

    def _toggle_word_wrap(self):
        """Toggle word wrap"""
        wrap_mode = "word" if self.word_wrap_var.get() else "none"
        self.text_widget.config(wrap=wrap_mode)

    def _change_font_size(self, value):
        """Change font size"""
        size = int(float(value))
        self.text_widget.config(font=("Consolas", size))

    def _update_text_display(self, text):
        """Update text display with optional line numbers"""
        self.text_widget.delete("1.0", "end")
        if self.line_numbers_var.get():
            lines = text.split('\n')
            numbered_text = ""
            for i, line in enumerate(lines, 1):
                numbered_text += f"{i:4d} | {line}\n"
            self.text_widget.insert("1.0", numbered_text)
        else:
            self.text_widget.insert("1.0", text)
        self.text_widget.config(state="disabled")

    def _show_image(self):
        """Enhanced image viewer with zoom and rotation"""
        self.original_image = Image.open(io.BytesIO(self.file_bytes))
        self.current_image = self.original_image.copy()
        self.zoom_factor = 1.0
        self.rotation_angle = 0
        
        # Image controls
        controls_frame = ctk.CTkFrame(self.control_frame)
        controls_frame.pack(fill="x", padx=5, pady=5)
        
        # Zoom controls
        ctk.CTkLabel(controls_frame, text="Zoom:").pack(side="left", padx=5, pady=5)
        zoom_out_btn = ctk.CTkButton(controls_frame, text="-", command=self._zoom_out, width=30)
        zoom_out_btn.pack(side="left", padx=2, pady=5)
        
        self.zoom_label = ctk.CTkLabel(controls_frame, text="100%", width=50)
        self.zoom_label.pack(side="left", padx=5, pady=5)
        
        zoom_in_btn = ctk.CTkButton(controls_frame, text="+", command=self._zoom_in, width=30)
        zoom_in_btn.pack(side="left", padx=2, pady=5)
        
        reset_btn = ctk.CTkButton(controls_frame, text="Reset", command=self._reset_image, width=60)
        reset_btn.pack(side="left", padx=10, pady=5)
        
        # Rotation controls
        ctk.CTkLabel(controls_frame, text="Rotate:").pack(side="left", padx=(20, 5), pady=5)
        rotate_left_btn = ctk.CTkButton(controls_frame, text="↺ 90°", command=self._rotate_left, width=60)
        rotate_left_btn.pack(side="left", padx=2, pady=5)
        
        rotate_right_btn = ctk.CTkButton(controls_frame, text="↻ 90°", command=self._rotate_right, width=60)
        rotate_right_btn.pack(side="left", padx=2, pady=5)
        
        # Image info
        info_text = f"Size: {self.original_image.size[0]}x{self.original_image.size[1]} | Mode: {self.original_image.mode}"
        ctk.CTkLabel(controls_frame, text=info_text).pack(side="right", padx=5, pady=5)
        
        # Image display
        self.image_label = ctk.CTkLabel(self.content_frame, text="")
        self.image_label.pack(expand=True, fill="both", padx=10, pady=10)
        
        self._update_image_display()

    def _zoom_in(self):
        """Zoom in the image"""
        self.zoom_factor = min(self.zoom_factor * 1.2, 5.0)
        self._update_image_display()

    def _zoom_out(self):
        """Zoom out the image"""
        self.zoom_factor = max(self.zoom_factor / 1.2, 0.1)
        self._update_image_display()

    def _rotate_left(self):
        """Rotate image 90° counter-clockwise"""
        self.rotation_angle = (self.rotation_angle - 90) % 360
        self._update_image_display()

    def _rotate_right(self):
        """Rotate image 90° clockwise"""
        self.rotation_angle = (self.rotation_angle + 90) % 360
        self._update_image_display()

    def _reset_image(self):
        """Reset image to original state"""
        self.zoom_factor = 1.0
        self.rotation_angle = 0
        self._update_image_display()

    def _update_image_display(self):
        """Update the image display with current transformations"""
        # Apply rotation
        if self.rotation_angle != 0:
            self.current_image = self.original_image.rotate(-self.rotation_angle, expand=True)
        else:
            self.current_image = self.original_image.copy()
        
        # Apply zoom
        new_size = (
            int(self.current_image.size[0] * self.zoom_factor),
            int(self.current_image.size[1] * self.zoom_factor)
        )
        
        # Limit size to reasonable bounds
        max_size = (1200, 800)
        if new_size[0] > max_size[0] or new_size[1] > max_size[1]:
            self.current_image.thumbnail(max_size, Image.Resampling.LANCZOS)
        else:
            self.current_image = self.current_image.resize(new_size, Image.Resampling.LANCZOS)
        
        # Update display
        photo = ImageTk.PhotoImage(self.current_image)
        self.image_label.configure(image=photo)
        self.image_label.image = photo
        
        # Update zoom label
        self.zoom_label.configure(text=f"{int(self.zoom_factor * 100)}%")

    def _show_pdf(self):
        """Enhanced PDF viewer with navigation"""
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(self.file_bytes)
            tmp_path = tmp.name
        
        try:
            self.pdf_doc = fitz.open(tmp_path)
            self.total_pdf_pages = self.pdf_doc.page_count
            self.current_pdf_page = 0
            
            # PDF controls
            controls_frame = ctk.CTkFrame(self.control_frame)
            controls_frame.pack(fill="x", padx=5, pady=5)
            
            # Navigation controls
            first_btn = ctk.CTkButton(controls_frame, text="⏮", command=self._pdf_first_page, width=40)
            first_btn.pack(side="left", padx=2, pady=5)
            
            prev_btn = ctk.CTkButton(controls_frame, text="⏪", command=self._pdf_prev_page, width=40)
            prev_btn.pack(side="left", padx=2, pady=5)
            
            self.pdf_page_label = ctk.CTkLabel(controls_frame, text="", width=100)
            self.pdf_page_label.pack(side="left", padx=10, pady=5)
            
            next_btn = ctk.CTkButton(controls_frame, text="⏩", command=self._pdf_next_page, width=40)
            next_btn.pack(side="left", padx=2, pady=5)
            
            last_btn = ctk.CTkButton(controls_frame, text="⏭", command=self._pdf_last_page, width=40)
            last_btn.pack(side="left", padx=2, pady=5)
            
            # Zoom controls for PDF
            ctk.CTkLabel(controls_frame, text="Zoom:").pack(side="left", padx=(20, 5), pady=5)
            self.pdf_zoom = ctk.DoubleVar(value=1.0)
            pdf_zoom_slider = ctk.CTkSlider(controls_frame, from_=0.5, to=3.0,
                                          variable=self.pdf_zoom, command=self._update_pdf_zoom, width=150)
            pdf_zoom_slider.pack(side="left", padx=5, pady=5)
            
            self.pdf_zoom_label = ctk.CTkLabel(controls_frame, text="100%", width=50)
            self.pdf_zoom_label.pack(side="left", padx=5, pady=5)
            
            # Page jump
            ctk.CTkLabel(controls_frame, text="Go to page:").pack(side="right", padx=5, pady=5)
            self.page_entry = ctk.CTkEntry(controls_frame, width=60)
            self.page_entry.pack(side="right", padx=5, pady=5)
            self.page_entry.bind("<Return>", self._jump_to_page)
            
            # PDF display
            self.pdf_label = ctk.CTkLabel(self.content_frame, text="")
            self.pdf_label.pack(expand=True, fill="both", padx=10, pady=10)
            
            self._update_pdf_display()
        
        except Exception as e:
            ctk.CTkLabel(self.content_frame, text=f"Failed to open PDF: {e}").pack(padx=20, pady=20)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def _pdf_first_page(self):
        """Go to first page"""
        self.current_pdf_page = 0
        self._update_pdf_display()

    def _pdf_last_page(self):
        """Go to last page"""
        self.current_pdf_page = self.total_pdf_pages - 1
        self._update_pdf_display()

    def _pdf_prev_page(self):
        """Go to previous page"""
        if self.current_pdf_page > 0:
            self.current_pdf_page -= 1
            self._update_pdf_display()

    def _pdf_next_page(self):
        """Go to next page"""
        if self.current_pdf_page < self.total_pdf_pages - 1:
            self.current_pdf_page += 1
            self._update_pdf_display()

    def _jump_to_page(self, event=None):
        """Jump to specific page"""
        try:
            page_num = int(self.page_entry.get()) - 1  # Convert to 0-based index
            if 0 <= page_num < self.total_pdf_pages:
                self.current_pdf_page = page_num
                self._update_pdf_display()
            else:
                messagebox.showerror("Error", f"Page must be between 1 and {self.total_pdf_pages}")
        except ValueError:
            messagebox.showerror("Error", "Please enter a valid page number")

    def _update_pdf_zoom(self, value):
        """Update PDF zoom"""
        self._update_pdf_display()
        self.pdf_zoom_label.configure(text=f"{int(float(value) * 100)}%")

    def _update_pdf_display(self):
        """Update PDF page display"""
        if self.pdf_doc and 0 <= self.current_pdf_page < self.total_pdf_pages:
            page = self.pdf_doc.load_page(self.current_pdf_page)
            mat = fitz.Matrix(self.pdf_zoom.get(), self.pdf_zoom.get())
            pix = page.get_pixmap(matrix=mat)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            
            # Fit to display area while maintaining aspect ratio
            display_size = (900, 600)
            img.thumbnail(display_size, Image.Resampling.LANCZOS)
            
            photo = ImageTk.PhotoImage(img)
            self.pdf_label.configure(image=photo)
            self.pdf_label.image = photo
            
            # Update page label
            self.pdf_page_label.configure(
                text=f"Page {self.current_pdf_page + 1} of {self.total_pdf_pages}"
            )

    def _show_video(self):
        """Enhanced video player with controls (video and audio playback)"""
        with tempfile.NamedTemporaryFile(delete=False, suffix=self.file_ext) as tmp:
            tmp.write(self.file_bytes)
            self.video_path = tmp.name

        self.video_cap = cv2.VideoCapture(self.video_path)
        if not self.video_cap.isOpened():
            ctk.CTkLabel(self.content_frame, text="Failed to open video.").pack(padx=20, pady=20)
            return

        # Get video properties
        self.video_fps = self.video_cap.get(cv2.CAP_PROP_FPS) or 30
        self.total_frames = int(self.video_cap.get(cv2.CAP_PROP_FRAME_COUNT))
        self.current_frame = 0

        # Initialize audio for video
        self.video_has_audio = False
        self.audio_error_message = None
        if self._initialize_audio():
            try:
                # Create temporary audio file
                with tempfile.NamedTemporaryFile(delete=False, suffix='.wav') as audio_tmp:
                    self.video_audio_path = audio_tmp.name
                # Try using ffmpeg to extract audio (if available)
                try:
                    subprocess.run([
                        'ffmpeg', '-i', self.video_path, 
                        '-vn', '-acodec', 'pcm_s16le', 
                        '-ar', '22050', '-ac', '2', 
                        self.video_audio_path
                    ], check=True, capture_output=True)
                    pygame.mixer.music.load(self.video_audio_path)
                    self.video_has_audio = True
                except (subprocess.CalledProcessError, FileNotFoundError) as e:
                    self.audio_error_message = ("Audio playback unavailable: ffmpeg is not installed or audio extraction failed. "
                                               "Install ffmpeg and ensure it is in your system PATH for video audio support.")
                    # Fallback: try loading video file directly
                    try:
                        pygame.mixer.music.load(self.video_path)
                        self.video_has_audio = True
                        self.audio_error_message = None
                    except Exception as e2:
                        self.audio_error_message = ("Audio playback unavailable: ffmpeg is not installed or audio extraction failed. "
                                                   "Install ffmpeg and ensure it is in your system PATH for video audio support.")
            except Exception as e:
                print(f"Video audio setup failed: {e}")
                self.video_has_audio = False
                self.audio_error_message = str(e)

        # Video controls
        controls_frame = ctk.CTkFrame(self.control_frame)
        controls_frame.pack(fill="x", padx=5, pady=5)

        # Playback controls
        self.play_btn = ctk.CTkButton(controls_frame, text="▶", command=self._toggle_video_play, width=40)
        self.play_btn.pack(side="left", padx=5, pady=5)

        stop_btn = ctk.CTkButton(controls_frame, text="⏹", command=self._stop_video, width=40)
        stop_btn.pack(side="left", padx=2, pady=5)

        # Progress slider
        self.video_progress = ctk.DoubleVar()
        self.progress_slider = ctk.CTkSlider(
            controls_frame, from_=0, to=self.total_frames-1,
            variable=self.video_progress, command=self._seek_video,
            width=300
        )
        self.progress_slider.pack(side="left", padx=10, pady=5, fill="x", expand=True)

        # Time display
        self.time_label = ctk.CTkLabel(controls_frame, text="00:00 / 00:00", width=100)
        self.time_label.pack(side="right", padx=5, pady=5)

        # Speed control
        ctk.CTkLabel(controls_frame, text="Speed:").pack(side="right", padx=5, pady=5)
        self.speed_var = ctk.DoubleVar(value=1.0)
        speed_slider = ctk.CTkSlider(controls_frame, from_=0.25, to=2.0,
                                   variable=self.speed_var, width=100)
        speed_slider.pack(side="right", padx=5, pady=5)

        # Video display
        self.video_label = ctk.CTkLabel(self.content_frame, text="")
        self.video_label.pack(expand=True, fill="both", padx=10, pady=10)

        # Show audio error message if needed
        if self.audio_error_message:
            ctk.CTkLabel(self.content_frame, text=self.audio_error_message, text_color="orange", font=ctk.CTkFont(size=12)).pack(pady=10)

        # Show first frame
        self._show_video_frame()

    def _toggle_video_play(self):
        """Toggle video play/pause"""
        if self.video_playing:
            self._pause_video()
        else:
            self._play_video()

    def _play_video(self):
        """Start/resume video playback (with audio)"""
        self.video_playing = True
        self.video_paused = False
        self.play_btn.configure(text="⏸")
        
        # Start audio if available
        if self.video_has_audio:
            try:
                if not pygame.mixer.music.get_busy():
                    pygame.mixer.music.play()
                elif self.video_paused:
                    pygame.mixer.music.unpause()
            except Exception as e:
                print(f"Video Audio Play error: {e}")

        if not self.video_thread or not self.video_thread.is_alive():
            self.video_thread = threading.Thread(target=self._video_playback_loop)
            self.video_thread.daemon = True
            self.video_thread.start()

    def _pause_video(self):
        """Pause video playback (and audio)"""
        self.video_playing = False
        self.video_paused = True
        self.play_btn.configure(text="▶")
        
        if self.video_has_audio:
            try:
                pygame.mixer.music.pause()
            except Exception as e:
                print(f"Video Audio Pause error: {e}")

    def _stop_video(self):
        """Stop video playback (and audio)"""
        self.video_playing = False
        self.video_paused = False
        self.current_frame = 0
        self.video_cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
        self.video_progress.set(0)
        self.play_btn.configure(text="▶")
        
        if self.video_has_audio:
            try:
                pygame.mixer.music.stop()
            except Exception as e:
                print(f"Video Audio Stop error: {e}")
        
        self._show_video_frame()

    def _seek_video(self, value):
        """Seek to specific frame"""
        frame_num = int(float(value))
        self.current_frame = frame_num
        self.video_cap.set(cv2.CAP_PROP_POS_FRAMES, frame_num)
        self._show_video_frame()

    def _video_playback_loop(self):
        """Video playback loop"""
        while self.video_playing and self.video_cap and self.video_cap.isOpened():
            if not self.video_paused:
                ret, frame = self.video_cap.read()
                if ret:
                    self.current_frame += 1
                    self.after(0, lambda f=frame: self._update_video_frame(f))
                    self.after(0, self._update_video_progress)
                    
                    # Control playback speed
                    delay = 1.0 / (self.video_fps * self.speed_var.get())
                    time.sleep(delay)
                else:
                    # End of video
                    self.after(0, self._stop_video)
                    break
            else:
                time.sleep(0.1)

    def _show_video_frame(self):
        """Show current video frame"""
        ret, frame = self.video_cap.read()
        if ret:
            self._update_video_frame(frame)

    def _update_video_frame(self, frame):
        """Update video frame display"""
        frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        img = Image.fromarray(frame_rgb)
        img.thumbnail((800, 600), Image.Resampling.LANCZOS)
        
        photo = ImageTk.PhotoImage(img)
        self.video_label.configure(image=photo)
        self.video_label.image = photo

    def _update_video_progress(self):
        """Update video progress display"""
        if self.total_frames > 0:
            self.video_progress.set(self.current_frame)
            
            # Update time display
            current_time = self.current_frame / self.video_fps
            total_time = self.total_frames / self.video_fps
            current_str = f"{int(current_time//60):02d}:{int(current_time%60):02d}"
            total_str = f"{int(total_time//60):02d}:{int(total_time%60):02d}"
            self.time_label.configure(text=f"{current_str} / {total_str}")

    def _show_audio(self):
        """Enhanced audio player with better compatibility"""
        with tempfile.NamedTemporaryFile(delete=False, suffix=self.file_ext) as tmp:
            tmp.write(self.file_bytes)
            self.audio_path = tmp.name

        # Initialize audio system
        if not self._initialize_audio():
            ctk.CTkLabel(self.content_frame, 
                        text="Failed to initialize audio system").pack(padx=20, pady=20)
            return

        try:
            pygame.mixer.music.load(self.audio_path)
            audio_loaded = True
        except Exception as e:
            # Try converting to a more compatible format
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix='.wav') as wav_tmp:
                    wav_path = wav_tmp.name
                
                subprocess.run([
                    'ffmpeg', '-i', self.audio_path,
                    '-acodec', 'pcm_s16le', '-ar', '22050', '-ac', '2',
                    wav_path
                ], check=True, capture_output=True)
                
                pygame.mixer.music.load(wav_path)
                self.audio_path = wav_path  # Update path to converted file
                audio_loaded = True
            except:
                audio_loaded = False
                
        if not audio_loaded:
            ctk.CTkLabel(self.content_frame, 
                        text=f"Unsupported audio format: {self.file_ext}").pack(padx=20, pady=20)
            return

        # Audio info
        info_frame = ctk.CTkFrame(self.content_frame)
        info_frame.pack(fill="x", padx=10, pady=10)

        info_text = f"🎵 Audio File: {self.filename}\n📁 Size: {len(self.file_bytes)} bytes"
        ctk.CTkLabel(info_frame, text=info_text, font=ctk.CTkFont(size=14)).pack(pady=20)

        # Audio controls
        controls_frame = ctk.CTkFrame(self.control_frame)
        controls_frame.pack(fill="x", padx=5, pady=5)

        # Playback controls
        self.audio_play_btn = ctk.CTkButton(controls_frame, text="▶",
                                          command=self._toggle_audio_play, width=50)
        self.audio_play_btn.pack(side="left", padx=5, pady=5)

        stop_btn = ctk.CTkButton(controls_frame, text="⏹",
                               command=self._stop_audio, width=50)
        stop_btn.pack(side="left", padx=2, pady=5)

        # Volume control
        ctk.CTkLabel(controls_frame, text="Volume:").pack(side="left", padx=(20, 5), pady=5)
        self.volume_var = ctk.DoubleVar(value=0.7)
        volume_slider = ctk.CTkSlider(controls_frame, from_=0, to=1,
                                    variable=self.volume_var,
                                    command=self._change_volume, width=150)
        volume_slider.pack(side="left", padx=5, pady=5)

        self.volume_label = ctk.CTkLabel(controls_frame, text="70%", width=40)
        self.volume_label.pack(side="left", padx=5, pady=5)

        # Status display
        self.audio_status_label = ctk.CTkLabel(controls_frame, text="Ready to play", width=120)
        self.audio_status_label.pack(side="right", padx=5, pady=5)

        # Set initial volume
        pygame.mixer.music.set_volume(0.7)

    def _toggle_audio_play(self):
        """Toggle audio play/pause"""
        if self.audio_playing:
            self._pause_audio()
        else:
            self._play_audio()

    def _play_audio(self):
        """Start/resume audio playback"""
        try:
            if self.audio_paused:
                pygame.mixer.music.unpause()
                self._show_audio_status("Resumed")
            else:
                pygame.mixer.music.play()
                self._show_audio_status("Playing...")

            self.audio_playing = True
            self.audio_paused = False
            self.audio_play_btn.configure(text="⏸")

            # Start monitoring playback
            self._monitor_audio_playback()

        except Exception as e:
            self._show_audio_status(f"Error: {e}", is_error=True)

    def _pause_audio(self):
        """Pause audio playback"""
        try:
            pygame.mixer.music.pause()
            self.audio_playing = False
            self.audio_paused = True
            self.audio_play_btn.configure(text="▶")
            self._show_audio_status("Paused")
        except Exception as e:
            self._show_audio_status(f"Error: {e}", is_error=True)

    def _stop_audio(self):
        """Stop audio playback"""
        try:
            pygame.mixer.music.stop()
            self.audio_playing = False
            self.audio_paused = False
            self.audio_play_btn.configure(text="▶")
            self._show_audio_status("Stopped")
        except Exception as e:
            self._show_audio_status(f"Error: {e}", is_error=True)

    def _change_volume(self, value):
        """Change audio volume"""
        volume = float(value)
        pygame.mixer.music.set_volume(volume)
        self.volume_label.configure(text=f"{int(volume * 100)}%")

    def _monitor_audio_playback(self):
        """Monitor audio playback status"""
        if self.audio_playing and pygame.mixer.music.get_busy():
            # Still playing, check again after 100ms
            self.after(100, self._monitor_audio_playback)
        elif self.audio_playing:
            # Playback finished
            self.audio_playing = False
            self.audio_paused = False
            self.audio_play_btn.configure(text="▶")
            self._show_audio_status("Finished")

    def _show_pptx(self):
        """Enhanced PowerPoint viewer with slide navigation"""
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp:
            tmp.write(self.file_bytes)
            tmp_path = tmp.name

        try:
            prs = Presentation(tmp_path)
            self.pptx_slides = []

            # Extract slide content
            for i, slide in enumerate(prs.slides):
                slide_content = f"=== Slide {i+1} ===\n\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_content += shape.text + "\n"
                    elif shape.shape_type == 13:  # Picture
                        slide_content += "[Image]\n"
                    elif shape.shape_type == 6:  # Table
                        slide_content += "[Table]\n"

                if not slide_content.strip().endswith("==="):
                    self.pptx_slides.append(slide_content)
                else:
                    self.pptx_slides.append(slide_content + "No content found on this slide.\n")

            if not self.pptx_slides:
                self.pptx_slides = ["No slides found in presentation."]

            self.current_pptx_slide = 0

            # PowerPoint controls
            controls_frame = ctk.CTkFrame(self.control_frame)
            controls_frame.pack(fill="x", padx=5, pady=5)

            prev_slide_btn = ctk.CTkButton(controls_frame, text="⏪ Previous",
                                         command=self._pptx_prev_slide, width=80)
            prev_slide_btn.pack(side="left", padx=5, pady=5)

            self.pptx_slide_label = ctk.CTkLabel(controls_frame, text="", width=120)
            self.pptx_slide_label.pack(side="left", padx=10, pady=5)

            next_slide_btn = ctk.CTkButton(controls_frame, text="Next ⏩",
                                         command=self._pptx_next_slide, width=80)
            next_slide_btn.pack(side="left", padx=5, pady=5)

            # Slide overview button
            overview_btn = ctk.CTkButton(controls_frame, text="Overview",
                                       command=self._show_pptx_overview, width=80)
            overview_btn.pack(side="right", padx=5, pady=5)

            # Text display for slides
            self.pptx_text_widget = scrolledtext.ScrolledText(
                self.content_frame,
                wrap="word",
                font=("Arial", 12),
                state="normal"
            )
            self.pptx_text_widget.pack(expand=True, fill="both", padx=10, pady=10)

            self._update_pptx_display()

        except Exception as e:
            ctk.CTkLabel(self.content_frame, text=f"Failed to display PPTX: {e}").pack(padx=20, pady=20)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def _pptx_prev_slide(self):
        """Show previous slide"""
        if self.current_pptx_slide > 0:
            self.current_pptx_slide -= 1
            self._update_pptx_display()

    def _pptx_next_slide(self):
        """Show next slide"""
        if self.current_pptx_slide < len(self.pptx_slides) - 1:
            self.current_pptx_slide += 1
            self._update_pptx_display()

    def _update_pptx_display(self):
        """Update PowerPoint slide display"""
        if 0 <= self.current_pptx_slide < len(self.pptx_slides):
            self.pptx_text_widget.config(state="normal")
            self.pptx_text_widget.delete("1.0", "end")
            self.pptx_text_widget.insert("1.0", self.pptx_slides[self.current_pptx_slide])
            self.pptx_text_widget.config(state="disabled")

            self.pptx_slide_label.configure(
                text=f"Slide {self.current_pptx_slide + 1} of {len(self.pptx_slides)}"
            )

    def _show_pptx_overview(self):
        """Show all slides overview"""
        overview_text = ""
        for i, slide in enumerate(self.pptx_slides):
            overview_text += f"{'='*50}\n{slide}\n"

        self.pptx_text_widget.config(state="normal")
        self.pptx_text_widget.delete("1.0", "end")
        self.pptx_text_widget.insert("1.0", overview_text)
        self.pptx_text_widget.config(state="disabled")
        self.pptx_slide_label.configure(text="Overview Mode")

    def _show_unsupported(self):
        """Show unsupported file message with file info"""
        info_frame = ctk.CTkFrame(self.content_frame)
        info_frame.pack(expand=True, fill="both", padx=20, pady=20)

        # File icon (using emoji)
        icon_map = {
            '.doc': '📄', '.docx': '📄', '.xls': '📊', '.xlsx': '📊',
            '.zip': '🗜️', '.rar': '🗜️', '.7z': '🗜️',
            '.exe': '⚙️', '.msi': '⚙️',
            '.iso': '💿', '.dmg': '💿'
        }

        icon = icon_map.get(self.file_ext, '📁')
        ctk.CTkLabel(info_frame, text=icon, font=ctk.CTkFont(size=48)).pack(pady=20)

        ctk.CTkLabel(info_frame, text=f"File Type: {self.file_ext.upper()}",
                    font=ctk.CTkFont(size=16, weight="bold")).pack(pady=5)

        ctk.CTkLabel(info_frame, text="This file type is not supported for preview.",
                    font=ctk.CTkFont(size=12)).pack(pady=5)

        # Show file size and other info
        size_mb = len(self.file_bytes) / (1024 * 1024)
        if size_mb > 1:
            size_text = f"Size: {size_mb:.2f} MB"
        else:
            size_kb = len(self.file_bytes) / 1024
            size_text = f"Size: {size_kb:.2f} KB"

        ctk.CTkLabel(info_frame, text=size_text).pack(pady=5)

        # Suggest required dependencies if applicable
        suggestions = {
            '.pdf': "Install PyMuPDF: pip install PyMuPDF",
            '.mp4': "Install OpenCV: pip install opencv-python",
            '.mp3': "Install Pygame: pip install pygame",
            '.pptx': "Install python-pptx: pip install python-pptx"
        }

        if self.file_ext in suggestions:
            ctk.CTkLabel(info_frame, text=f"💡 Tip: {suggestions[self.file_ext]}",
                        text_color="orange").pack(pady=10)

    def on_closing(self):
        """Clean up resources when closing"""
        try:
            # Stop all audio/video playback
            if hasattr(self, 'video_playing') and self.video_playing:
                self.video_playing = False
            if hasattr(self, 'audio_playing') and self.audio_playing:
                self.audio_playing = False
            
            # Stop pygame mixer
            if pygame and pygame.mixer.get_init():
                pygame.mixer.music.stop()
                pygame.mixer.quit()
            
            # Clean up video resources
            if hasattr(self, 'video_cap') and self.video_cap:
                self.video_cap.release()
            
            # Clean up PDF
            if hasattr(self, 'pdf_doc') and self.pdf_doc:
                self.pdf_doc.close()
            
            # Remove temporary files
            temp_files = ['video_path', 'audio_path', 'video_audio_path']
            for attr in temp_files:
                if hasattr(self, attr):
                    path = getattr(self, attr)
                    if path and os.path.exists(path):
                        os.remove(path)
                        
        except Exception as e:
            print(f"Cleanup error: {e}")
        finally:
            self.destroy()


# Example usage
if __name__ == "__main__":
    ctk.set_appearance_mode("dark")
    ctk.set_default_color_theme("blue")
    
    app = ctk.CTk()
    app.withdraw()  # Hide main window for this example
    
    # Example with a text file
    sample_text = "Hello World!\nThis is a sample text file.\n" * 100
    viewer = FileViewer(app, sample_text.encode(), "sample.txt")
    
    app.mainloop()
